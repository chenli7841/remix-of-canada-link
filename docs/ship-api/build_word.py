from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'word'
OUT.mkdir(exist_ok=True)

def font(style, size=10.5, mono=False):
    style.font.name = 'Consolas' if mono else 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')

def inline(p, text):
    for i, part in enumerate(re.split(r'(`[^`]+`)', text)):
        r=p.add_run(part[1:-1] if part.startswith('`') else part)
        if part.startswith('`'):
            r.font.name='Consolas'
            r.font.size=Pt(9)

def convert(source, filename):
    d=Document()
    sec=d.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(1.9)
    sec.left_margin=sec.right_margin=Cm(1.8)
    for s in ['Normal','List Bullet','List Number']:
        font(d.styles[s]); d.styles[s].paragraph_format.space_after=Pt(5)
        d.styles[s].paragraph_format.line_spacing=1.16
    for s,sz in [('Title',22),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
        font(d.styles[s],sz)
        d.styles[s].paragraph_format.space_before=Pt(12)
        d.styles[s].paragraph_format.space_after=Pt(7)
    code=d.styles.add_style('Code Sample',1)
    font(code,8.5,True)
    code.paragraph_format.space_after=Pt(0)
    code.paragraph_format.line_spacing=1.0
    foot=sec.footer.paragraphs[0]; foot.alignment=2
    r=foot.add_run(); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); r._r.addnext(f)
    lines=source.read_text(encoding='utf-8-sig').splitlines()
    i=0; incode=False
    while i<len(lines):
        line=lines[i]
        if line.startswith(('~~~','```')):
            incode=not incode; i+=1; continue
        if incode:
            p=d.add_paragraph(style='Code Sample'); p.add_run(line)
            i+=1; continue
        if not line.strip(): i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                row=[s.strip() for s in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',v) for v in row): rows.append(row)
                i+=1
            n=len(rows[0]); t=d.add_table(rows=0, cols=n)
            t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
            ratios=([.40,.60] if n==2 else [.16,.46,.38] if rows[0][0]=='方法' else [.26,.34,.40])
            if len(ratios)!=n: ratios=[1/n]*n
            for col,ratio in zip(t.columns,ratios): col.width=Cm(17.4*ratio)
            pr=t._tbl.tblPr
            borders=OxmlElement('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            for ri,row in enumerate(rows):
                rr=t.add_row()
                if ri==0:
                    header=OxmlElement('w:tblHeader'); rr._tr.get_or_add_trPr().append(header)
                for ci,value in enumerate(row):
                    c=rr.cells[ci]; c.width=Cm(17.4*ratios[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=c.paragraphs[0]; inline(p,value)
                    p.paragraph_format.space_after=Pt(4); p.paragraph_format.space_before=Pt(4)
                    p.paragraph_format.line_spacing=1.08
                    for r in p.runs:
                        r.font.size=Pt(9); r.bold=ri==0
                    tc=c._tc.get_or_add_tcPr()
                    margins=OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+side); e.set(qn('w:w'),'80'); e.set(qn('w:type'),'dxa'); margins.append(e)
                    tc.append(margins)
                    if ri==0:
                        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE6F1'); tc.append(shade)
            d.add_paragraph().paragraph_format.space_after=Pt(2)
            continue
        m=re.match(r'^(#{1,4})\s+(.*)',line)
        if m:
            level=len(m[1]); title=re.sub(r'[：:（）()、/，。]+',' ',m[2])
            title=re.sub(r'^(\d+)\.\s*',r'\1 ',title)
            p=d.add_paragraph(style='Title' if level==1 else f'Heading {level-1}')
            p.add_run(title)
        elif line.startswith('- '): inline(d.add_paragraph(style='List Bullet'),line[2:])
        elif re.match(r'^\d+\. ',line): inline(d.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',line))
        else: inline(d.add_paragraph(),line)
        i+=1
    d.core_properties.title=d.paragraphs[0].text
    d.core_properties.author=''
    path=OUT/filename; d.save(path)
    reopened=Document(path)
    assert len(reopened.paragraphs)>50
    print(f'{path} | paragraphs={len(reopened.paragraphs)} tables={len(reopened.tables)}')

convert(ROOT/'system-change-guide-v2.md','系统修改意见与Claude指导_v2.docx')
convert(ROOT/'shipper-api-v2.md','Shipper_API对接文档_v2.docx')

